from __future__ import annotations

import argparse
import json
import os
import random
import re
import shutil
import subprocess
import sys
import textwrap
import time
from functools import lru_cache
from pathlib import Path
from typing import Callable, Iterable

import requests
from PIL import Image, ImageDraw, ImageFilter, ImageFont, features

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
except Exception:  # pragma: no cover - graceful fallback for minimal installs
    arabic_reshaper = None
    get_display = None

try:
    import uharfbuzz as hb
    import freetype
    HARFBUZZ_OK = True
except Exception:  # pragma: no cover - graceful fallback for minimal installs
    hb = None
    freetype = None
    HARFBUZZ_OK = False

try:
    import imageio_ffmpeg
except Exception:  # pragma: no cover
    imageio_ffmpeg = None

# When frozen into a single EXE (PyInstaller), __file__ points inside the
# temporary extraction folder that gets wiped between runs — user data
# (output, temp, backgrounds, the task database) must live next to the EXE
# itself instead, so it survives across launches.
BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
BUNDLE_DIR = Path(getattr(sys, "_MEIPASS", BASE_DIR))
TEMP_DIR = BASE_DIR / "temp"
OUTPUT_DIR = BASE_DIR / "output"
BACKGROUND_DIR = BASE_DIR / "خلفيات"
FONT_DIR = BASE_DIR / "fonts"
TASKS_XLSX = BASE_DIR / "tasks.xlsx"
# Where each background was left off, so a long clip keeps feeding new videos
# from where the last one stopped instead of replaying its opening every time.
# It lives next to the EXE, not in temp/, because temp gets cleared.
BG_STATE_FILE = BASE_DIR / "bg_state.json"

RECITERS = {
    "ياسر الدوسري": {"src": "m", "id": "4", "every_id": "Yasser_Ad-Dussary_128kbps"},
    "مشاري راشد العفاسي": {"src": "m", "id": "1", "every_id": "Alafasy_128kbps"},
    "ناصر القطامي": {"src": "m", "id": "3", "every_id": "Nasser_Alqatami_128kbps"},
    "أبو بكر الشاطري": {"src": "m", "id": "2", "every_id": "Abu_Bakr_Ash-Shaatree_128kbps"},
    "محمد صديق المنشاوي — مرتّل": {"src": "e", "every_id": "Minshawy_Murattal_128kbps"},
    "محمد صديق المنشاوي — مجوّد": {"src": "e", "every_id": "Minshawy_Mujawwad_192kbps"},
    "عبد الباسط عبد الصمد — مرتّل": {"src": "e", "every_id": "Abdul_Basit_Murattal_192kbps"},
    "عبد الباسط عبد الصمد — مجوّد": {"src": "e", "every_id": "Abdul_Basit_Mujawwad_128kbps"},
    "ماهر المعيقلي": {"src": "e", "every_id": "MaherAlMuaiqly128kbps"},
    "محمود خليل الحصري — مرتّل": {"src": "e", "every_id": "Husary_128kbps"},
    "محمود خليل الحصري — مجوّد": {"src": "e", "every_id": "Husary_128kbps_Mujawwad"},
    "سعود الشريم": {"src": "e", "every_id": "Saood_ash-Shuraym_128kbps"},
    "عبد الرحمن السديس": {"src": "e", "every_id": "Abdurrahmaan_As-Sudais_192kbps"},
    "محمود علي البنا": {"src": "e", "every_id": "mahmoud_ali_al_banna_32kbps"},
    "هاني الرفاعي": {"src": "m", "id": "5", "every_id": "Hani_Rifai_192kbps"},
}

ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".avi"}
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}

# Selectable body fonts, mirroring the reference builder's font menu. Each
# maps to a file in fonts/; missing files fall back to the Quran font.
FONT_CHOICES = {
    "AmiriQuran": ("أميري قرآن", "AmiriQuran.ttf"),
    "Amiri": ("أميري", "Amiri-Bold.ttf"),
    "ScheherazadeNew": ("شهرزاد", "ScheherazadeNew-Regular.ttf"),
    "NotoNaskhArabic": ("نوتو نسخ", "NotoNaskhArabic-Regular.ttf"),
}

FONT_DOWNLOADS = {
    "ScheherazadeNew-Regular.ttf": "https://github.com/google/fonts/raw/main/ofl/scheherazadenew/ScheherazadeNew-Regular.ttf",
    "NotoNaskhArabic-Regular.ttf": "https://github.com/google/fonts/raw/main/ofl/notonaskharabic/NotoNaskhArabic%5Bwght%5D.ttf",
}

# One place defining every appearance setting, so the control panel, the
# preview and every clip in a batch all read the same defaults.
DEFAULT_STYLE = {
    "font": "AmiriQuran",
    "font_size": 62,
    "text_color": "#FFFFFF",
    "position": "center",
    "tashkeel": True,
    "show_text": True,
    "show_number": True,
    "stroke": True,
    "text_bg": False,
    "text_bg_color": "#000000",
    "text_bg_style": "box",
    "text_bg_opacity": 45,
    "text_bg_pad": 0.55,
    "text_bg_radius": 0.30,
    "shadow": True,
    "shadow_opacity": 55,
}

# How the background is fitted into the 9:16 frame. "cover" enlarges it until
# it fills the frame and crops the overflow; "blur" keeps the whole picture at
# its own aspect ratio over a blurred copy of itself, so nothing is enlarged
# and nothing is cut. These two constants keep the FFmpeg branch and the
# Pillow preview branch producing the same look.
BLUR_FILL_SIGMA = 28
BLUR_FILL_DARKEN = 0.08

# Arabic tafsir editions offered in the control panel, keyed by the id used by
# api.quran.com. المیسر is first because it is the short, plain-language one
# that suits a video description.
TAFSIR_CHOICES = {
    "muyassar": (16, "التفسير الميسر", "ar-tafsir-muyassar"),
    "saadi": (91, "تفسير السعدي", "ar-tafseer-al-saddi"),
    "ibnkathir": (14, "تفسير ابن كثير", "ar-tafsir-ibn-kathir"),
    "tabari": (15, "تفسير الطبري", "ar-tafsir-al-tabari"),
    "qurtubi": (90, "تفسير القرطبي", "ar-tafseer-al-qurtubi"),
    "baghawi": (94, "تفسير البغوي", "ar-tafsir-al-baghawi"),
    "waseet": (93, "التفسير الوسيط", "ar-tafsir-al-wasit"),
}


def log(message: str, callback: Callable[[str], None] | None = None) -> None:
    stamp = time.strftime("%H:%M:%S")
    line = f"[{stamp}] {message}"
    print(line, flush=True)
    if callback:
        callback(line)


def executable(name: str) -> str | None:
    found = shutil.which(name)
    if found:
        return found
    if name == "ffmpeg" and imageio_ffmpeg:
        try:
            return imageio_ffmpeg.get_ffmpeg_exe()
        except Exception:
            pass
    return None


def run_ffmpeg(args: list[str], *, callback=None) -> subprocess.CompletedProcess:
    ffmpeg = executable("ffmpeg")
    if not ffmpeg:
        raise RuntimeError("لم يتم العثور على FFmpeg. ثبّته أو أعد تشغيل install_windows.bat.")
    command = [ffmpeg, "-hide_banner", "-loglevel", "error", *args]
    result = subprocess.run(command, capture_output=True, text=True)
    if result.returncode != 0:
        detail = (result.stderr or result.stdout or "خطأ غير معروف في FFmpeg").strip()
        raise RuntimeError(detail[-1800:])
    return result


def text_kwargs() -> dict:
    """Use native RTL shaping when Pillow is built with libraqm."""
    try:
        if features.check("raqm"):
            return {"direction": "rtl", "language": "ar"}
    except Exception:
        pass
    return {}


_fallback_reshaper = None
if arabic_reshaper:
    try:
        # arabic_reshaper's default configuration deletes tashkeel
        # (delete_harakat=True) because plain presentation-form glyphs
        # can't carry combining marks safely. That default silently
        # stripped every diacritic from ayah text on installs without
        # libraqm, so it must be turned off explicitly here.
        _fallback_reshaper = arabic_reshaper.ArabicReshaper(configuration={
            "delete_harakat": False,
            "shift_harakat_position": True,
            "support_ligatures": True,
        })
    except Exception:
        _fallback_reshaper = None


def prepare_arabic(value: str) -> str:
    # Keep the API string unchanged when native shaping is available.
    value = str(value or "")
    if HARFBUZZ_OK or text_kwargs():
        return value
    if _fallback_reshaper and get_display:
        try:
            return get_display(_fallback_reshaper.reshape(value))
        except Exception:
            pass
    return value


def arabic_number(value: int | str) -> str:
    return str(value).translate(ARABIC_DIGITS)


def safe_name(value: str) -> str:
    value = re.sub(r"[^\w\-\u0600-\u06ff ]+", "", str(value), flags=re.UNICODE)
    return re.sub(r"\s+", "_", value.strip())[:100] or "clip"


def load_tasks(path: Path = TASKS_XLSX) -> list[dict]:
    from openpyxl import load_workbook

    if not path.exists() or path.stat().st_size == 0:
        bundled = BUNDLE_DIR / path.name
        if bundled.exists() and bundled.stat().st_size > 0:
            path = bundled
        else:
            raise FileNotFoundError(f"لم يتم العثور على قاعدة البيانات أو أنها فارغة: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    header_index = next((i for i, row in enumerate(rows) if row and "المعرّف" in [str(x) for x in row]), None)
    if header_index is None:
        raise ValueError("تعذر العثور على صف عناوين قاعدة البيانات")
    headers = [str(x).strip() if x is not None else "" for x in rows[header_index]]
    tasks: list[dict] = []
    for row in rows[header_index + 1 :]:
        if not row or not row[1]:
            continue
        record = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        tasks.append({
            "index": int(record.get("م") or len(tasks) + 1),
            "id": str(record.get("المعرّف") or f"Q-{len(tasks)+1:03d}"),
            "juz": int(record.get("الجزء") or 0),
            "surah_number": int(record.get("رقم السورة") or 0),
            "surah_name": str(record.get("اسم السورة") or ""),
            "from_ayah": int(record.get("من آية") or 1),
            "to_ayah": int(record.get("إلى آية") or 1),
            "ayah_count": int(record.get("عدد الآيات") or 0),
            "opening": str(record.get("مطلع الآية الأولى") or ""),
            "reciter": str(record.get("القارئ") or "ياسر الدوسري"),
            "status": "pending",
        })
    return tasks


def _extract_ayahs(payload) -> dict[int, str]:
    """Extract the Arabic text exactly as supplied by the selected API payload."""
    result: dict[int, str] = {}
    if isinstance(payload, dict):
        # quranapi.pages.dev returns the API-owned Arabic text in arabic1.
        # Prefer it over any translated or locally reconstructed text.
        for key in ("arabic1", "arabic_uthmani", "arabic"):
            arabic_values = payload.get(key)
            if isinstance(arabic_values, list) and arabic_values:
                return {index + 1: str(text) for index, text in enumerate(arabic_values) if text is not None}
        candidates = payload.get("data", payload)
        if isinstance(candidates, dict):
            candidates = candidates.get("ayahs", candidates.get("verses", candidates.get("Ayahs", candidates)))
    else:
        candidates = payload
    if isinstance(candidates, dict):
        if "text" in candidates:
            number = candidates.get("numberInSurah", candidates.get("number", 1))
            return {int(number): str(candidates["text"])}
        items = candidates.values()
    else:
        items = candidates if isinstance(candidates, list) else []
    for item in items:
        if not isinstance(item, dict):
            continue
        text = item.get("text") or item.get("text_uthmani") or item.get("content")
        number = item.get("numberInSurah") or item.get("ayah") or item.get("number")
        if text is not None and number is not None:
            try:
                result[int(number)] = str(text)
            except (TypeError, ValueError):
                continue
    return result


def fetch_surah_text(surah_number: int, *, session: requests.Session | None = None, timeout: int = 25) -> dict[int, dict[str, str]]:
    """Return {ayah_number: {"tashkeel": ..., "plain": ...}} for a surah.

    quranapi.pages.dev serves the fully-vocalised text in `arabic1` and the
    undotted/plain variant in `arabic2`; both are kept so the diacritics can be
    toggled at render time without a second network round-trip.
    """
    session = session or requests.Session()
    cache = TEMP_DIR / "text_cache" / f"surah_{surah_number:03d}_v2.json"
    if cache.exists():
        try:
            data = json.loads(cache.read_text(encoding="utf-8"))
            return {int(k): {"tashkeel": str(v["tashkeel"]), "plain": str(v["plain"])} for k, v in data.items()}
        except Exception:
            cache.unlink(missing_ok=True)

    verses: dict[int, dict[str, str]] = {}
    errors = []
    try:
        response = session.get(f"https://quranapi.pages.dev/api/{surah_number}.json", timeout=timeout)
        response.raise_for_status()
        payload = response.json()
        vocalised = payload.get("arabic1") or []
        plain = payload.get("arabic2") or []
        total = int(payload.get("totalAyah") or len(vocalised))
        if vocalised and len(vocalised) == total:
            for position, text in enumerate(vocalised, start=1):
                verses[position] = {
                    "tashkeel": str(text),
                    "plain": str(plain[position - 1]) if position - 1 < len(plain) else str(text),
                }
        elif vocalised:
            errors.append(
                f"quranapi: عدد الآيات المُعاد ({len(vocalised)}) لا يطابق العدد المعلن ({total})"
            )
    except Exception as exc:
        errors.append(f"quranapi: {exc}")

    if not verses:
        # Fallback keeps each ayah keyed by its own numberInSurah rather than
        # by list position, so a gap in the response can never shift the
        # remaining ayat onto the wrong numbers.
        try:
            response = session.get(
                f"https://api.alquran.cloud/v1/surah/{surah_number}/quran-uthmani", timeout=timeout
            )
            response.raise_for_status()
            for item in response.json().get("data", {}).get("ayahs", []) or []:
                number = item.get("numberInSurah")
                text = item.get("text")
                if number is None or text is None:
                    continue
                verses[int(number)] = {"tashkeel": str(text), "plain": str(text)}
        except Exception as exc:
            errors.append(f"alquran.cloud: {exc}")

    if not verses:
        raise RuntimeError("تعذر تحميل نصوص السورة من المصادر المتاحة. " + " | ".join(errors))

    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(verses, ensure_ascii=False, indent=2), encoding="utf-8")
    return verses


def get_ayah_texts(surah_number: int, start: int, end: int, *, tashkeel: bool = True, session=None) -> list[tuple[int, str]]:
    verses = fetch_surah_text(surah_number, session=session)
    missing = [str(n) for n in range(start, end + 1) if n not in verses]
    if missing:
        raise RuntimeError(f"الآيات التالية غير موجودة في المصدر: {', '.join(missing)}")
    key = "tashkeel" if tashkeel else "plain"
    return [(n, verses[n][key]) for n in range(start, end + 1)]


def audio_urls(surah: int, ayah: int, reciter_name: str) -> list[str]:
    info = RECITERS.get(reciter_name) or RECITERS["ياسر الدوسري"]
    urls: list[str] = []
    if info.get("src") == "m" and info.get("id"):
        rid = info["id"]
        urls.extend([
            f"https://the-quran-project.github.io/Quran-Audio/Data/{rid}/{surah}_{ayah}.mp3",
            f"https://the-quran-project.github.io/Quran-Audio/Data/{rid}/{surah:03d}{ayah:03d}.mp3",
        ])
    folder = info.get("every_id")
    if folder:
        urls.append(f"https://everyayah.com/data/{folder}/{surah:03d}{ayah:03d}.mp3")
    return urls


def download_audio(surah: int, ayah: int, reciter_name: str, *, session=None, callback=None) -> Path:
    session = session or requests.Session()
    target_dir = TEMP_DIR / "audio" / safe_name(reciter_name)
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{surah:03d}{ayah:03d}.mp3"
    if target.exists() and target.stat().st_size > 1000:
        return target
    last_error = ""
    for url in audio_urls(surah, ayah, reciter_name):
        try:
            response = session.get(url, timeout=45, stream=True)
            if response.status_code != 200:
                last_error = f"HTTP {response.status_code}"
                continue
            temp = target.with_suffix(".part")
            with temp.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=1024 * 128):
                    if chunk:
                        handle.write(chunk)
            if temp.stat().st_size > 1000:
                temp.replace(target)
                return target
            temp.unlink(missing_ok=True)
            last_error = "ملف صوتي فارغ"
        except Exception as exc:
            last_error = str(exc)
    raise RuntimeError(f"تعذر تحميل صوت سورة {surah} آية {ayah}: {last_error}")


def probe_duration(path: Path) -> float:
    ffprobe = executable("ffprobe")
    if ffprobe:
        result = subprocess.run(
            [ffprobe, "-v", "error", "-show_entries", "format=duration", "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
            capture_output=True, text=True,
        )
        try:
            return max(0.1, float(result.stdout.strip()))
        except (TypeError, ValueError):
            pass
    ffmpeg = executable("ffmpeg")
    if ffmpeg:
        result = subprocess.run([ffmpeg, "-i", str(path)], capture_output=True, text=True)
        match = re.search(r"Duration:\s+(\d+):(\d+):(\d+(?:\.\d+)?)", result.stderr)
        if match:
            h, m, s = match.groups()
            return max(0.1, int(h) * 3600 + int(m) * 60 + float(s))
    return 4.0


@lru_cache(maxsize=32)
def probe_size(path: Path) -> tuple[int, int]:
    """Pixel size of a background's first video stream, or (0, 0) if unknown.

    Cached because every ayah's segment asks for the same background."""
    ffprobe = executable("ffprobe")
    if ffprobe:
        result = subprocess.run(
            [ffprobe, "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=width,height",
             "-of", "csv=p=0:s=x", str(path)],
            capture_output=True, text=True,
        )
        match = re.search(r"(\d{2,5})x(\d{2,5})", result.stdout or "")
        if match:
            return int(match.group(1)), int(match.group(2))
    ffmpeg = executable("ffmpeg")
    if ffmpeg:
        # No ffprobe in the bundled imageio-ffmpeg build, so read the size out
        # of the banner FFmpeg prints when asked to open a file with no output.
        result = subprocess.run([ffmpeg, "-i", str(path)], capture_output=True, text=True)
        match = re.search(r"Video:.*?[\s,](\d{2,5})x(\d{2,5})", result.stderr or "")
        if match:
            return int(match.group(1)), int(match.group(2))
    return 0, 0


def cover_scale(width: int, height: int, *, target_w: int = 1080, target_h: int = 1920) -> float:
    """How much a background gets enlarged to fill the output frame. Above 1 it
    is being blown up, and the result needs sharpening to stay presentable."""
    if width <= 0 or height <= 0:
        return 1.0
    return max(target_w / width, target_h / height)


def merge_audio(audio_files: list[Path], target: Path) -> None:
    concat_file = target.with_suffix(".concat.txt")
    lines = []
    for file in audio_files:
        escaped = str(file.resolve()).replace("'", "'\\''")
        lines.append(f"file '{escaped}'")
    concat_file.write_text("\n".join(lines) + "\n", encoding="utf-8")
    try:
        run_ffmpeg(["-y", "-f", "concat", "-safe", "0", "-i", str(concat_file), "-c", "copy", str(target)])
    finally:
        concat_file.unlink(missing_ok=True)


def ensure_font() -> Path | None:
    FONT_DIR.mkdir(parents=True, exist_ok=True)
    local = FONT_DIR / "Amiri-Bold.ttf"
    if local.exists() and local.stat().st_size > 10000:
        return local
    bundled = BUNDLE_DIR / "fonts" / "Amiri-Bold.ttf"
    if bundled.exists() and bundled.stat().st_size > 10000:
        return bundled
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/trado.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    try:
        url = "https://raw.githubusercontent.com/alif-type/amiri/master/fonts/Amiri-Bold.ttf"
        response = requests.get(url, timeout=20)
        if response.status_code == 200 and len(response.content) > 10000:
            local.write_bytes(response.content)
            return local
    except Exception:
        pass
    return None


def ensure_quran_font() -> Path | None:
    """Return the dedicated Quran font (AmiriQuran carries the Uthmani-style
    diacritic anchors), shaped ourselves via HarfBuzz so it no longer needs
    Pillow's libraqm — see shape_arabic() below."""
    local = FONT_DIR / "AmiriQuran.ttf"
    if local.exists() and local.stat().st_size > 10000:
        return local
    bundled = BUNDLE_DIR / "fonts" / "AmiriQuran.ttf"
    if bundled.exists() and bundled.stat().st_size > 10000:
        return bundled
    return ensure_font()


def font_choice_path(name: str) -> Path | None:
    """Resolve a font menu key to a usable file, downloading it once if needed."""
    entry = FONT_CHOICES.get(str(name))
    if not entry:
        return None
    filename = entry[1]
    for candidate in (FONT_DIR / filename, BUNDLE_DIR / "fonts" / filename):
        if candidate.exists() and candidate.stat().st_size > 10000:
            return candidate
    url = FONT_DOWNLOADS.get(filename)
    if url:
        try:
            response = requests.get(url, timeout=25)
            if response.status_code == 200 and len(response.content) > 10000:
                FONT_DIR.mkdir(parents=True, exist_ok=True)
                target = FONT_DIR / filename
                target.write_bytes(response.content)
                return target
        except Exception:
            pass
    return None


def font(size: int, path: Path | None = None):
    try:
        if path:
            return ImageFont.truetype(str(path), size)
    except Exception:
        pass
    return ImageFont.load_default()


# ---------------------------------------------------------------------------
# HarfBuzz + FreeType Arabic shaping.
#
# Windows Pillow wheels are built without libraqm, so Pillow's own draw.text()
# cannot perform OpenType mark-attachment (GPOS "mark"/"mkmk") — every
# Quranic harakah would either be dropped (arabic_reshaper's default
# behaviour) or drawn floating in the wrong place. uharfbuzz + freetype-py
# ship prebuilt wheels on every platform, so shaping and rasterising glyphs
# by hand here gives correct tashkeel placement without depending on the
# host Pillow build at all.
# ---------------------------------------------------------------------------

_hb_font_cache: dict[tuple[str, int], "hb.Font"] = {}
_ft_face_cache: dict[str, "freetype.Face"] = {}


def _shapers_for(font_path: Path, size_px: int):
    key = (str(font_path), size_px)
    cached = _hb_font_cache.get(key)
    if cached:
        # A FreeType face carries one active pixel size and is cached per file,
        # shared by every size. Drawing the ayah-number digits re-scales that
        # shared face, so the size must be re-asserted on each use or the next
        # line of body text rasterises at the digits' size while HarfBuzz still
        # positions it for the body size.
        cached[1].set_char_size(size_px * 64)
        return cached
    face_ft = _ft_face_cache.get(str(font_path))
    if face_ft is None:
        face_ft = freetype.Face(str(font_path))
        _ft_face_cache[str(font_path)] = face_ft
    face_ft.set_char_size(size_px * 64)
    blob = hb.Blob.from_file_path(str(font_path))
    font_hb = hb.Font(hb.Face(blob))
    font_hb.scale = (size_px * 64, size_px * 64)
    hb.ot_font_set_funcs(font_hb)
    _hb_font_cache[key] = (font_hb, face_ft)
    return font_hb, face_ft


def shape_arabic(text: str, font_path: Path, size_px: int):
    """Shape text right-to-left and return (glyphs, total_width, face).

    Each glyph is (glyph_id, pen_x, y_offset) in pixels, already in final
    left-to-right drawing order for the shaped run.
    """
    font_hb, face_ft = _shapers_for(font_path, size_px)
    buf = hb.Buffer()
    buf.add_str(text)
    buf.guess_segment_properties()
    buf.direction = "rtl"
    buf.script = "Arab"
    buf.language = "ar"
    hb.shape(font_hb, buf)
    glyphs = []
    pen_x = 0.0
    for info, pos in zip(buf.glyph_infos, buf.glyph_positions):
        glyphs.append((info.codepoint, pen_x + pos.x_offset / 64, pos.y_offset / 64))
        pen_x += pos.x_advance / 64
    return glyphs, pen_x, face_ft


def measure_arabic(text: str, font_path: Path, size_px: int) -> float:
    if not text:
        return 0.0
    _, width, _ = shape_arabic(text, font_path, size_px)
    return width


def _glyph_ink_bounds(glyphs, face_ft) -> tuple[float, float]:
    top, bottom = None, None
    for codepoint, _gx, gy_off in glyphs:
        face_ft.load_glyph(codepoint, freetype.FT_LOAD_RENDER)
        g = face_ft.glyph
        if g.bitmap.rows <= 0:
            continue
        glyph_top = -gy_off - g.bitmap_top
        glyph_bottom = glyph_top + g.bitmap.rows
        top = glyph_top if top is None else min(top, glyph_top)
        bottom = glyph_bottom if bottom is None else max(bottom, glyph_bottom)
    if top is None:
        return 0.0, 0.0
    return top, bottom


def draw_arabic(card: Image.Image, text: str, font_path: Path, size_px: int, x: float, y: float, *,
                 valign: str = "middle", fill=(255, 255, 255, 255),
                 stroke_width: int = 0, stroke_fill=(0, 0, 0, 200)) -> list[tuple[float, float, float]]:
    """Draw one line of shaped, centered Arabic text onto card at (x, y).

    Returns the placements of any END OF AYAH ornaments drawn, as
    (center_x, center_y, width) — the caller draws the ayah digits inside them.
    """
    if not text:
        return []
    glyphs, total_width, face_ft = shape_arabic(text, font_path, size_px)
    if not glyphs:
        return []
    start_x = x - total_width / 2
    if valign == "middle":
        top, bottom = _glyph_ink_bounds(glyphs, face_ft)
        base_y = y - (top + bottom) / 2
    else:  # "top": y marks the top of the ink
        top, _bottom = _glyph_ink_bounds(glyphs, face_ft)
        base_y = y - top

    marker_gid = face_ft.get_char_index(0x06DD)
    markers: list[tuple[float, float, float]] = []
    rendered = []
    for codepoint, gx_off, gy_off in glyphs:
        face_ft.load_glyph(codepoint, freetype.FT_LOAD_RENDER)
        g = face_ft.glyph
        bmp = g.bitmap
        if bmp.width <= 0 or bmp.rows <= 0:
            continue
        mask = Image.frombytes("L", (bmp.width, bmp.rows), bytes(bmp.buffer))
        gx = start_x + gx_off + g.bitmap_left
        gy = base_y - gy_off - g.bitmap_top
        if codepoint == marker_gid:
            # Callers append a bare U+06DD and let draw_ayah_number() composite
            # the digits inside it. Shaping "۝" together with its digits makes
            # HarfBuzz report their offsets against the ornament's RTL (right
            # edge) origin, which places them outside the ornament entirely.
            markers.append((gx + bmp.width / 2, gy + bmp.rows / 2, float(bmp.width)))
        rendered.append((mask, gx, gy))

    if stroke_width > 0:
        for mask, gx, gy in rendered:
            for dx in range(-stroke_width, stroke_width + 1):
                for dy in range(-stroke_width, stroke_width + 1):
                    if dx == 0 and dy == 0:
                        continue
                    card.paste(Image.new("RGBA", mask.size, stroke_fill), (int(gx + dx), int(gy + dy)), mask)
    for mask, gx, gy in rendered:
        card.paste(Image.new("RGBA", mask.size, fill), (int(gx), int(gy)), mask)
    return markers


def draw_ayah_number(card: Image.Image, number: int, font_path: Path, marker, *,
                     fill=(255, 255, 255, 255)) -> None:
    """Draw the ayah digits centered inside an already-drawn ۝ ornament."""
    center_x, center_y, marker_width = marker
    digits = arabic_number(number)
    # Sized to sit inside the ornament's inner counter with a little air.
    size_px = max(9, int(marker_width * (0.46 if len(digits) < 3 else 0.36)))
    glyphs, width, face_ft = shape_arabic(digits, font_path, size_px)
    if not glyphs:
        return
    top, bottom = _glyph_ink_bounds(glyphs, face_ft)
    start_x = center_x - width / 2
    base_y = center_y - (top + bottom) / 2
    for codepoint, gx_off, gy_off in glyphs:
        face_ft.load_glyph(codepoint, freetype.FT_LOAD_RENDER)
        g = face_ft.glyph
        bmp = g.bitmap
        if bmp.width <= 0 or bmp.rows <= 0:
            continue
        mask = Image.frombytes("L", (bmp.width, bmp.rows), bytes(bmp.buffer))
        card.paste(
            Image.new("RGBA", mask.size, fill),
            (int(start_x + gx_off + g.bitmap_left), int(base_y - gy_off - g.bitmap_top)),
            mask,
        )


def wrap_arabic(text: str, font_path: Path | None, max_width: float, size: int) -> list[str]:
    """Greedy word wrap at a fixed size, mirroring the reference wrapLines()."""
    words = str(text or "").split()
    use_hb = HARFBUZZ_OK and font_path is not None
    draw = None if use_hb else ImageDraw.Draw(Image.new("RGBA", (1, 1)))

    def width_of(candidate: str) -> float:
        if use_hb:
            return measure_arabic(candidate, font_path, size)
        bbox = draw.textbbox((0, 0), prepare_arabic(candidate), font=font(size, font_path), **text_kwargs())
        return bbox[2] - bbox[0]

    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if not current or width_of(trial) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines if use_hb else [prepare_arabic(x) for x in lines]


def fit_wrapped(text: str, font_path: Path | None, max_width: int, max_lines: int, start_size: int) -> tuple[int, list[str]]:
    """Word-wrap Arabic text and return (font_size_px, lines)."""
    words = str(text or "").split()
    use_hb = HARFBUZZ_OK and font_path is not None

    def width_of(candidate: str, size: int) -> float:
        if use_hb:
            return measure_arabic(candidate, font_path, size)
        fnt = font(size, font_path)
        draw = ImageDraw.Draw(Image.new("RGBA", (1, 1)))
        bbox = draw.textbbox((0, 0), prepare_arabic(candidate), font=fnt, **text_kwargs())
        return bbox[2] - bbox[0]

    for size in range(start_size, 24, -2):
        lines: list[str] = []
        current = ""
        for word in words:
            trial = f"{current} {word}".strip()
            if width_of(trial, size) <= max_width or not current:
                current = trial
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
        if len(lines) <= max_lines:
            return size, (lines if use_hb else [prepare_arabic(x) for x in lines])
    fallback_text = " ".join(words[:80])
    return 24, ([fallback_text] if use_hb else [prepare_arabic(fallback_text)])


def hex_rgba(value: str, alpha: float) -> tuple[int, int, int, int]:
    text = str(value or "#000000").lstrip("#")
    if len(text) == 3:
        text = "".join(c * 2 for c in text)
    try:
        r, g, b = int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)
    except ValueError:
        r, g, b = 0, 0, 0
    return (r, g, b, int(255 * max(0.0, min(1.0, alpha))))


# Layout constants mirroring the reference builder's buildLayout()/drawText().
LINE_HEIGHT_RATIO = 1.62
WRAP_WIDTH_RATIO = 0.84
STROKE_RATIO = 0.085
BG_PAD_Y_RATIO = 0.76
POSITION_ANCHORS = {"top": 0.16, "center": 0.50, "bottom": 0.84}
# Reference canvas used shadowBlur = size*0.22 and shadowOffsetY = size*0.04.
# Canvas shadowBlur is twice the Gaussian sigma, hence the /2 when blurring.
SHADOW_BLUR_RATIO = 0.22
SHADOW_OFFSET_RATIO = 0.04
SHADOW_GAIN = 3.0


def render_card(*, ayah_number: int, ayah_text: str, output: Path,
                width: int = 1080, height: int = 1920,
                font_path: Path | None = None, quran_font_path: Path | None = None,
                show_text: bool = True, show_number: bool = True,
                font_size: int = 62, text_color: str = "#FFFFFF",
                position: str = "center", stroke: bool = True,
                text_bg: bool = False, text_bg_color: str = "#000000",
                text_bg_style: str = "box", text_bg_opacity: int = 45,
                text_bg_pad: float = 0.55, text_bg_radius: float = 0.30,
                shadow: bool = True, shadow_opacity: int = 55,
                bg_dim: int = 0) -> None:
    """Render one ayah exactly as the reference canvas builder lays it out:
    the ayah text alone, its number inline as ۝, on a transparent card."""
    card = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(card)
    body_font_path = quran_font_path or font_path
    # The background dim rides on this overlay rather than being a separate
    # FFmpeg blend stage: FFmpeg would mix it in YUV (limited range) while the
    # preview composites in RGB, and the two came out visibly different.
    # Baking it here means both paths run the identical alpha composite.
    dim_alpha = int(255 * max(0, min(100, bg_dim)) / 100)
    if dim_alpha > 0:
        draw.rectangle((0, 0, width, height), fill=(0, 0, 0, dim_alpha))
    if not show_text or not body_font_path:
        output.parent.mkdir(parents=True, exist_ok=True)
        card.save(output, "PNG")
        return

    # Font size scales with the canvas, exactly as the reference does with
    # size * (W / 1080), so a setting looks identical at any resolution.
    size = max(12, int(font_size * (width / 1080)))
    text = str(ayah_text or "")
    if show_number:
        # The bare ornament joins the text run; its digits are composited
        # inside it afterwards by draw_ayah_number().
        text = f"{text} ۝"

    lines = wrap_arabic(text, body_font_path, width * WRAP_WIDTH_RATIO, size)
    if not lines:
        output.parent.mkdir(parents=True, exist_ok=True)
        card.save(output, "PNG")
        return

    use_hb = HARFBUZZ_OK and body_font_path is not None
    if use_hb:
        widths = [measure_arabic(line, body_font_path, size) for line in lines]
    else:
        widths = [draw.textbbox((0, 0), line, font=font(size, body_font_path), **text_kwargs())[2] for line in lines]

    line_height = size * LINE_HEIGHT_RATIO
    block = len(lines) * line_height
    anchor = POSITION_ANCHORS.get(position, 0.50)
    center_y = height * anchor + (block / 2 if position == "top" else -block / 2 if position == "bottom" else 0)
    y0 = center_y - block / 2 + line_height / 2

    if text_bg and text_bg_opacity > 0:
        fill = hex_rgba(text_bg_color, text_bg_opacity / 100)
        pad_x = size * text_bg_pad
        pad_y = size * text_bg_pad * BG_PAD_Y_RATIO
        radius = size * text_bg_radius
        if text_bg_style == "bar":
            draw.rectangle((0, y0 - line_height / 2 - pad_y, width, y0 - line_height / 2 + block + pad_y), fill=fill)
        elif text_bg_style == "lines":
            for index in range(len(lines)):
                top = y0 + index * line_height - line_height / 2 - pad_y * 0.3
                draw.rounded_rectangle(
                    (width / 2 - widths[index] / 2 - pad_x * 0.75, top,
                     width / 2 + widths[index] / 2 + pad_x * 0.75, top + line_height + pad_y * 0.6),
                    radius=radius, fill=fill)
        else:  # "box"
            max_width = max(widths)
            draw.rounded_rectangle(
                (width / 2 - max_width / 2 - pad_x, y0 - line_height / 2 - pad_y,
                 width / 2 + max_width / 2 + pad_x, y0 - line_height / 2 + block + pad_y),
                radius=radius, fill=fill)

    color = hex_rgba(text_color, 1.0)
    stroke_width = max(1, int(size * STROKE_RATIO / 2)) if stroke else 0
    # Text goes on its own layer so the soft drop shadow can be derived from
    # the finished glyph alpha and laid down underneath it, the way the
    # reference canvas applies shadowBlur/shadowOffsetY to each text draw.
    text_layer = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    layer_draw = ImageDraw.Draw(text_layer)
    for index, line in enumerate(lines):
        y = y0 + index * line_height
        if use_hb:
            markers = draw_arabic(text_layer, line, body_font_path, size, width / 2, y, valign="middle",
                                  fill=color, stroke_width=stroke_width, stroke_fill=(0, 0, 0, 158))
            for marker in markers:
                draw_ayah_number(text_layer, ayah_number, body_font_path, marker, fill=color)
        else:
            layer_draw.text((width / 2, y), line, font=font(size, body_font_path), fill=color, anchor="mm",
                            align="center", stroke_width=stroke_width, stroke_fill=(0, 0, 0, 158), **text_kwargs())

    if shadow and shadow_opacity > 0:
        sigma = max(1.0, size * SHADOW_BLUR_RATIO / 2)
        offset_y = int(round(size * SHADOW_OFFSET_RATIO))
        blurred = text_layer.getchannel("A").filter(ImageFilter.GaussianBlur(sigma))
        strength = max(0, min(100, shadow_opacity)) / 100
        # Blurring spreads a glyph's alpha thin, so scaling it by the strength
        # alone leaves even 100% barely visible. Applying gain and clamping
        # keeps the soft falloff while letting the slider reach a halo that
        # actually separates the text from a busy or bright background.
        gain = strength * SHADOW_GAIN
        blurred = blurred.point(lambda v, k=gain: min(255, int(v * k)))
        shadow_alpha = Image.new("L", (width, height), 0)
        shadow_alpha.paste(blurred, (0, offset_y))
        shadow_layer = Image.new("RGBA", (width, height), (0, 0, 0, 0))
        shadow_layer.putalpha(shadow_alpha)
        card.alpha_composite(shadow_layer)

    card.alpha_composite(text_layer)
    output.parent.mkdir(parents=True, exist_ok=True)
    card.save(output, "PNG")


def default_background(width: int = 1080, height: int = 1920) -> Image.Image:
    """Generated deep-navy gradient used whenever no background file exists,
    so previews and clips always show the same thing rather than plain black."""
    image = Image.new("RGB", (width, height), (10, 15, 29))
    draw = ImageDraw.Draw(image)
    for y in range(height):
        ratio = y / max(height - 1, 1)
        draw.line((0, y, width, y),
                  fill=(int(10 + 22 * ratio), int(15 + 20 * ratio), int(29 + 34 * ratio)))
    return image.filter(ImageFilter.GaussianBlur(0.6))


def background_still(background: Path | None, *, width: int = 1080, height: int = 1920,
                     dim: int = 0, bg_fit: str = "cover") -> Image.Image:
    """One still frame of a background, fitted to the output size and dimmed —
    the exact still the video pipeline produces for its first frame."""
    base: Image.Image | None = None
    if background and background.exists():
        try:
            if background.suffix.lower() in IMAGE_EXTS:
                base = Image.open(background).convert("RGB")
            else:
                frame = TEMP_DIR / "bg_frame.png"
                frame.parent.mkdir(parents=True, exist_ok=True)
                run_ffmpeg(["-y", "-ss", "0", "-i", str(background), "-frames:v", "1", str(frame)])
                base = Image.open(frame).convert("RGB")
        except Exception:
            base = None
    if base is None:
        base = default_background(width, height)
    else:
        # Same geometry as the FFmpeg filter: scale to cover, then centre-crop.
        scale = max(width / base.width, height / base.height)
        resized = base.resize((max(1, round(base.width * scale)), max(1, round(base.height * scale))), Image.LANCZOS)
        left = (resized.width - width) // 2
        top = (resized.height - height) // 2
        cover = resized.crop((left, top, left + width, top + height))
        if bg_fit == "blur":
            # Mirror the FFmpeg blur-fill branch: the whole frame shrunk to fit
            # inside, centred over a blurred, darkened cover-crop of itself. A
            # 9:16 background fits exactly and hides the blurred layer entirely.
            fit = min(width / base.width, height / base.height)
            inner = base.resize((max(1, round(base.width * fit)), max(1, round(base.height * fit))), Image.LANCZOS)
            canvas = cover.filter(ImageFilter.GaussianBlur(BLUR_FILL_SIGMA))
            canvas = Image.blend(canvas, Image.new("RGB", canvas.size, (0, 0, 0)), BLUR_FILL_DARKEN)
            canvas.paste(inner, ((width - inner.width) // 2, (height - inner.height) // 2))
            base = canvas
        else:
            base = cover
    amount = max(0, min(100, dim)) / 100
    if amount > 0:
        base = Image.blend(base, Image.new("RGB", base.size, (0, 0, 0)), amount)
    return base


def background_files() -> list[Path]:
    BACKGROUND_DIR.mkdir(parents=True, exist_ok=True)
    return sorted([p for p in BACKGROUND_DIR.iterdir() if p.is_file() and p.suffix.lower() in VIDEO_EXTS | IMAGE_EXTS])


# ---------------------------------------------------------------------------
# Background continuation
#
# A single ten-minute background can carry a dozen fifty-second videos without
# ever showing the same footage twice — but only if the next video knows where
# the last one stopped. These helpers keep one cursor per background file in
# bg_state.json: how many seconds of it have already been used.
# ---------------------------------------------------------------------------

# A background is treated as spent once less than this many seconds remain;
# a two-second tail is not worth starting a video on.
BG_MIN_REMAINING = 5.0


def bg_min_remaining(duration: float) -> float:
    """How short a tail makes a background not worth starting on.

    A flat five seconds was written for backgrounds minutes long, and it made
    short clips unusable: a ten-second stock shot dropped out of the pool after
    a single ayah, because five of its ten seconds always counted as too small
    a remainder to begin a video on. Scaling the threshold with the clip keeps
    the judgement sane at both ends."""
    if duration <= 0:
        return BG_MIN_REMAINING
    return min(BG_MIN_REMAINING, duration * 0.3)


def load_bg_state() -> dict:
    try:
        data = json.loads(BG_STATE_FILE.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        # A corrupt or missing state file must never stop a render — the worst
        # case of starting over is a repeated background, not a failure.
        return {}


def save_bg_state(state: dict) -> None:
    try:
        BG_STATE_FILE.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def background_fingerprint(path: Path) -> str:
    """Size and modification time of the file. Stored beside the cursor so that
    replacing a background with a different clip of the same name resets it
    instead of resuming at an offset that means nothing in the new footage."""
    try:
        stat = path.stat()
        return f"{stat.st_size}:{int(stat.st_mtime)}"
    except OSError:
        return ""


def background_usage(path: Path) -> dict:
    """Cursor, total length and remaining seconds for one background.

    The measured length is written back into the state file, because probing
    costs an FFmpeg launch per file and the control panel asks for the whole
    folder every time it refreshes."""
    state = load_bg_state()
    entry = state.get(path.name) or {}
    is_video = path.suffix.lower() in VIDEO_EXTS
    fingerprint = background_fingerprint(path)
    fresh = entry.get("fingerprint") == fingerprint
    cursor = float(entry.get("cursor", 0.0) or 0.0) if fresh else 0.0
    duration = float(entry.get("duration", 0.0) or 0.0) if fresh else 0.0
    if is_video and duration <= 0:
        duration = probe_duration(path)
        state[path.name] = {"cursor": round(cursor, 3), "fingerprint": fingerprint,
                            "duration": round(duration, 3)}
        save_bg_state(state)
    if is_video:
        cursor = min(max(cursor, 0.0), duration)
        remaining = max(0.0, duration - cursor)
    else:
        # A still has no timeline, so it is simply "used" or "unused" and takes
        # its turn in the rotation once per pass.
        remaining = 0.0 if cursor > 0 else 1.0
    return {"name": path.name, "cursor": cursor, "duration": duration,
            "remaining": remaining, "is_video": is_video}


def advance_background_cursor(path: Path, consumed: float) -> None:
    """Record that `consumed` seconds of this background have now been used."""
    usage = background_usage(path)
    state = load_bg_state()
    cursor = usage["cursor"] + consumed if usage["is_video"] else 1.0
    if usage["is_video"] and usage["duration"] > 0:
        cursor %= usage["duration"]
    state[path.name] = {"cursor": round(cursor, 3), "fingerprint": background_fingerprint(path),
                        "duration": round(usage["duration"], 3)}
    save_bg_state(state)


def reset_background_cursors() -> None:
    save_bg_state({})


def pick_background(choices: list[Path]) -> Path:
    """Random pick, but only among backgrounds that still have unseen footage.

    Plain random re-showed the opening of the same clip again and again. Drawing
    from the unspent ones instead walks through everything in the folder before
    anything repeats; when they are all spent the cursors reset and it starts a
    fresh pass."""
    fresh = [p for p in choices if background_has_footage(background_usage(p))]
    if not fresh:
        reset_background_cursors()
        fresh = choices
    return random.choice(fresh)


def background_has_footage(usage: dict) -> bool:
    """Whether this background still has unseen material worth starting on."""
    if not usage["is_video"]:
        return usage["remaining"] > 0
    return usage["remaining"] > bg_min_remaining(usage["duration"])


# ---------------------------------------------------------------------------
# Background planning
#
# A background long enough to carry the whole recitation is used the way it
# always was: one continuous take, resumed at the cursor the previous video
# left behind. Short clips cannot do that. A ten-second stock shot stretched
# over a ninety-second recitation wraps eight times, and because the wrap is a
# hard cut the viewer watches the same three seconds restart again and again.
#
# So when — and only when — the footage is too short for one take, the plan
# hands each ayah its own background. The picture then changes on the ayah
# boundary, where a cut lands with the recitation instead of across it.
# ---------------------------------------------------------------------------

BG_MODES = ("auto", "single", "per_ayah")

# Two backgrounds whose spare footage differs by less than this are treated as
# equally good fits, so the choice between them stays random instead of always
# landing on the same file.
BG_FIT_SLACK = 5.0


def usable_footage(usage: dict, bg_continue: bool) -> float:
    """Seconds this background can play before wrapping back to its start.

    With continuation on, only the part after the cursor is unseen; with it off
    the clip plays from the top and its whole length is available."""
    if not usage["is_video"]:
        return 0.0
    return usage["remaining"] if bg_continue else usage["duration"]


def single_background_plan(background: Path, durations: list[float], *, bg_continue: bool = True,
                           usage: dict | None = None) -> list[dict]:
    """One background across every ayah, each segment resuming where the last
    one stopped so the finished clip plays as a single unbroken take."""
    is_video = background.suffix.lower() in VIDEO_EXTS
    if usage is None and is_video:
        usage = background_usage(background)
    duration = float(usage["duration"]) if usage and is_video else 0.0
    start = float(usage["cursor"]) if (usage and is_video and bg_continue) else 0.0
    plan: list[dict] = []
    elapsed = 0.0
    for length in durations:
        offset = (start + elapsed) % duration if duration > 0.1 else 0.0
        plan.append({"path": background, "offset": offset})
        elapsed += length
    return plan


def plan_backgrounds(choices: list[Path], durations: list[float], *, bg_continue: bool = True,
                     mode: str = "auto", callback=None) -> tuple[list[dict], str]:
    """Decide which background carries each ayah, and from which offset.

    Returns the per-ayah plan and the mode it settled on, so the caller can
    report which of the two it actually used."""
    if mode not in BG_MODES:
        mode = "auto"
    total = sum(durations)
    usages = {p: background_usage(p) for p in choices}
    if bg_continue and choices and not any(background_has_footage(usages[p]) for p in choices):
        # Every background has been played to its end. The single-clip picker
        # already handled this by wiping the cursors and starting a fresh pass;
        # without the same reset here the planner would hand each ayah the last
        # fraction of a second of some clip and wrap immediately — the worst
        # possible output, and silent.
        reset_background_cursors()
        usages = {p: background_usage(p) for p in choices}
        log("انتهت كل الخلفيات المتاحة — إعادة الدورة من بدايتها", callback)

    local: dict[Path, float] = {}   # seconds this clip has already taken from each

    def footage(path: Path) -> float:
        return usable_footage(usages[path], bg_continue)

    def available(path: Path) -> float:
        """Unseen footage left on this background: the part the shared cursor
        has not reached, minus whatever this clip has already spent of it."""
        return footage(path) - local.get(path, 0.0)

    def covers(path: Path, needed: float) -> bool:
        # A still image never runs out, so it covers an ayah of any length.
        return True if not usages[path]["is_video"] else available(path) >= needed

    if mode == "auto" and not any(usages[p]["is_video"] for p in choices):
        # A folder of stills has no looping problem to solve, so auto leaves it
        # on the behaviour it already had. Per-ayah stays reachable by choosing
        # it explicitly, which turns such a folder into a slideshow.
        mode = "single"

    if mode == "single":
        # Prefer a background that can carry the whole clip without wrapping,
        # and fall back to the old pool when none can.
        pool = [p for p in choices if footage(p) >= total]
        if not pool:
            pool = [p for p in choices if background_has_footage(usages[p])] or list(choices)
        chosen = random.choice(pool)
        usage = usages[chosen]
        if bg_continue and usage["is_video"] and usage["cursor"] > 0:
            log(f"متابعة الخلفية «{chosen.name}» من {usage['cursor'] / 60:.0f}:"
                f"{usage['cursor'] % 60:04.1f}", callback)
        return single_background_plan(chosen, durations, bg_continue=bg_continue, usage=usage), "single"

    plan: list[dict] = []
    spent: list[Path] = []          # backgrounds this clip is finished with
    current: Path | None = None
    for needed in durations:
        # Hold the shot. A thirty-second background handed to a four-second
        # ayah has twenty-six seconds of unseen footage left, and dropping it
        # there throws that footage away for this video and cuts a long take
        # off mid-scene. It keeps carrying ayat until it genuinely runs out —
        # which for a short clip is the very next ayah, so those still change
        # every time, and for a background longer than the whole recitation is
        # never, so that case stays the single unbroken take it always was.
        if mode == "auto" and current is not None and usages[current]["is_video"] \
                and covers(current, needed):
            chosen = current
        else:
            if current is not None:
                spent.append(current)
            pool = [p for p in choices if p not in spent]
            if not pool:
                # Every background has been drawn on once already, so start a
                # second pass — but never straight back onto the one on screen,
                # which would read as the picture stalling rather than a cut.
                spent = [current] if current is not None else []
                pool = [p for p in choices if p not in spent] or list(choices)
            fits = [p for p in pool if covers(p, needed)]
            if fits:
                # Best fit, not a free-for-all. Handing a thirty-second clip to
                # a five-second ayah burns the only background long enough for
                # the twenty-five-second ayah further down the same video, and
                # that ayah is then stuck looping a ten-second clip. Taking the
                # tightest clip that still covers this ayah keeps the long ones
                # for the ayat that actually need them.
                surplus = {p: (float("inf") if not usages[p]["is_video"] else available(p) - needed)
                           for p in fits}
                tightest = min(surplus.values())
                chosen = random.choice([p for p in fits if surplus[p] <= tightest + BG_FIT_SLACK])
            else:
                # Nothing is long enough for this ayah, so the background with
                # the most footage left after its cursor wraps the fewest times.
                chosen = max(pool, key=available)
            current = chosen
        usage = usages[chosen]
        duration = usage["duration"] if usage["is_video"] else 0.0
        base = usage["cursor"] if (bg_continue and usage["is_video"]) else 0.0
        offset = (base + local.get(chosen, 0.0)) % duration if duration > 0.1 else 0.0
        plan.append({"path": chosen, "offset": offset})
        local[chosen] = local.get(chosen, 0.0) + needed
        if mode != "auto":
            # The explicit per-ayah mode never holds a shot: the user asked for
            # a change on every ayah, so retire each background immediately.
            spent.append(chosen)
            current = chosen
    distinct = len({entry["path"] for entry in plan})
    if distinct == 1:
        usage = usages[plan[0]["path"]]
        if bg_continue and usage["is_video"] and usage["cursor"] > 0:
            log(f"متابعة الخلفية «{plan[0]['path'].name}» من {usage['cursor'] / 60:.0f}:"
                f"{usage['cursor'] % 60:04.1f}", callback)
        return plan, "single"
    log(f"توزيع الخلفيات: {distinct} خلفية على {len(plan)} آية — "
        f"كل خلفية تُعرض حتى تنتهي ثم تنتقل للتالية", callback)
    return plan, "per_ayah"


def make_segment(background: Path | None, audio: Path, card: Path, output: Path, duration: float,
                 animation: str = "instant", bg_offset: float = 0.0, bg_fit: str = "cover") -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    args: list[str] = ["-y"]
    if background is None:
        args += ["-f", "lavfi", "-i", "color=c=#0A0F1D:s=1080x1920:r=30"]
    elif background.suffix.lower() in IMAGE_EXTS:
        args += ["-loop", "1", "-i", str(background)]
    else:
        # Each ayah is encoded as its own segment, so without a seek every
        # segment would start the background video at 00:00 and the finished
        # clip would visibly rewind on every ayah. `-ss` picks up the video
        # where the previous segment stopped; with `-stream_loop` the seek
        # applies to the first pass only, so once the background runs out it
        # wraps to its own start exactly as continuous playback would.
        args += ["-stream_loop", "-1"]
        if bg_offset > 0:
            args += ["-ss", f"{bg_offset:.3f}"]
        args += ["-i", str(background)]
    args += ["-i", str(audio), "-loop", "1", "-i", str(card)]
    fade = ",fade=t=in:st=0:d=0.8:alpha=1" if animation == "fade" else ""
    # The dim is already baked into the card's alpha by render_card().
    #
    # Quality notes:
    # * lanczos resampling — a landscape background has to be blown up a lot to
    #   fill a 9:16 frame (a 1280x720 clip is enlarged 2.7x), and the default
    #   bicubic kernel turns that enlargement to mush.
    # * a light unsharp afterwards restores most of the detail the enlargement
    #   costs; it is skipped when the picture is being shrunk instead, where
    #   sharpening would only add halos.
    # * overlay in RGB, then a single conversion to yuv420p at the very end.
    #   Blending the text card straight onto a subsampled 4:2:0 background
    #   averages the chroma of thin strokes with the background, which fringes
    #   the tashkeel; blending at full chroma resolution keeps them clean.
    size = probe_size(background) if background is not None else (0, 0)
    if bg_fit == "blur":
        # Nothing is cropped: the whole picture is shrunk to fit inside the
        # frame and centred over a blurred, darkened cover-crop of itself. A
        # background that is already 9:16 fits exactly and hides the blur.
        fit_scale = min(1080 / size[0], 1920 / size[1]) if size[0] > 0 and size[1] > 0 else 1.0
        sharpen = ",unsharp=5:5:0.8:5:5:0.0" if fit_scale > 1.15 else ""
        common = f"setsar=1,fps=30,trim=duration={duration:.3f}"
        filter_complex = (
            f"[0:v]split=2[blurbase][sharp];"
            f"[blurbase]scale=1080:1920:force_original_aspect_ratio=increase:flags=lanczos,crop=1080:1920,"
            f"gblur=sigma={BLUR_FILL_SIGMA},eq=brightness=-{BLUR_FILL_DARKEN},{common}[back];"
            f"[sharp]scale=1080:1920:force_original_aspect_ratio=decrease:flags=lanczos{sharpen},{common}[front];"
            f"[back][front]overlay=(W-w)/2:(H-h)/2[bg];"
            f"[2:v]format=rgba{fade}[card];[bg][card]overlay=0:0:format=rgb,format=yuv420p[v]"
        )
    else:
        sharpen = ",unsharp=5:5:0.8:5:5:0.0" if cover_scale(*size) > 1.15 else ""
        filter_complex = (
            f"[0:v]scale=1080:1920:force_original_aspect_ratio=increase:flags=lanczos,"
            f"crop=1080:1920{sharpen},setsar=1,fps=30,trim=duration={duration:.3f}[bg];"
            f"[2:v]format=rgba{fade}[card];[bg][card]overlay=0:0:format=rgb,format=yuv420p[v]"
        )
    args += ["-filter_complex", filter_complex, "-map", "[v]", "-map", "1:a:0", "-t", f"{duration:.3f}",
             "-r", "30", "-c:v", "libx264", "-preset", "medium", "-crf", "17", "-pix_fmt", "yuv420p",
             # Tag the colour explicitly; without it the filter chain leaves a
             # mismatched primaries/transfer/matrix triple on the file and some
             # players shift the colours.
             "-colorspace", "bt709", "-color_primaries", "bt709", "-color_trc", "bt709",
             "-c:a", "aac", "-b:a", "192k", "-movflags", "+faststart", str(output)]
    run_ffmpeg(args)


def concatenate_segments(segments: list[Path], output: Path) -> None:
    concat = output.with_suffix(".segments.txt")
    concat.write_text("\n".join(f"file '{str(p.resolve()).replace(chr(39), chr(39)+chr(92)+chr(39)+chr(39))}'" for p in segments) + "\n", encoding="utf-8")
    try:
        # +faststart moves the moov index to the front. Without it the
        # concatenated file keeps its index at the end, so a browser has to
        # download the whole clip before it can begin playing — which shows up
        # as a video that spins on "loading" forever in the gallery.
        run_ffmpeg(["-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-c", "copy",
                    "-movflags", "+faststart", str(output)])
    finally:
        concat.unlink(missing_ok=True)


def strip_html(value: str) -> str:
    """Some tafsir editions ship HTML (<p>, <br>, footnote spans). The caption
    file is plain text, so the markup has to come out."""
    text = re.sub(r"<br\s*/?>", "\n", value, flags=re.IGNORECASE)
    text = re.sub(r"</(p|div|h\d)>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    text = (text.replace("&nbsp;", " ").replace("&amp;", "&")
                .replace("&quot;", '"').replace("&lt;", "<").replace("&gt;", ">"))
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def get_tafsir(surah: int, ayah: int, edition: str = "muyassar",
               session: requests.Session | None = None) -> str:
    """Arabic tafsir for one ayah.

    Two independent sources are tried so a single outage cannot break a batch:
    the official quran.com v4 API first, then the same editions served as
    static JSON from the jsDelivr CDN. Returns "" when both fail — a missing
    tafsir must never abort a video that is otherwise fine.
    """
    tafsir_id, _, slug = TAFSIR_CHOICES.get(edition, TAFSIR_CHOICES["muyassar"])
    get = (session or requests).get
    sources = [
        (f"https://api.quran.com/api/v4/tafsirs/{tafsir_id}/by_ayah/{surah}:{ayah}",
         lambda payload: payload.get("tafsir", {}).get("text", "")),
        (f"https://cdn.jsdelivr.net/gh/spa5k/tafsir_api@main/tafsir/{slug}/{surah}/{ayah}.json",
         lambda payload: payload.get("text", "")),
    ]
    for url, extract in sources:
        try:
            response = get(url, timeout=20)
            response.raise_for_status()
            text = strip_html(str(extract(response.json()) or ""))
            if text:
                return text
        except Exception:
            continue
    return ""


def write_social_file(output_video: Path, task: dict, reciter: str, topic: str,
                      tafsir: list[tuple[int, str]] | None = None,
                      tafsir_name: str = "") -> Path:
    surah = task.get("surah_name", "سورة القرآن")
    start, end = task.get("from_ayah", 1), task.get("to_ayah", 1)
    tag_surah = re.sub(r"^سورة\s*", "", surah).replace(" ", "_")
    tag_reciter = re.sub(r"[^\u0600-\u06ffA-Za-z0-9]+", "_", reciter).strip("_")
    text = (
        f"تلاوة خاشعة تريح القلوب | {surah} ({start}-{end})\n"
        f"بصوت القارئ الشيخ {reciter}\n"
        f"الموضوع: {topic or 'تلاوة قرآنية خاشعة'}\n"
    )
    if tafsir:
        text += f"\n— {tafsir_name or 'التفسير'} —\n"
        for ayah_no, explanation in tafsir:
            text += f"\n﴿{ayah_no}﴾ {explanation}\n"
        text += "\n"
    text += f"هاشتاقات: #قرآن #{tag_surah} #{tag_reciter} #تلاوة_خاشعة #quran #fyp #shorts #reels\n"
    caption = output_video.with_name(output_video.stem + "_social.txt")
    caption.write_text(text, encoding="utf-8")
    return caption


def generate_task(task: dict, *, reciter: str | None = None, background: Path | None = None,
                  animation: str = "instant", style: dict | None = None, bg_dim: int = 0,
                  topic: str = "", bg_fit: str = "cover", tafsir: str = "",
                  bg_continue: bool = True, bg_mode: str = "auto",
                  callback=None, stop_check: Callable[[], bool] | None = None) -> Path:
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    style = dict(DEFAULT_STYLE, **(style or {}))
    session = requests.Session()
    reciter = reciter or task.get("reciter") or "ياسر الدوسري"
    surah_number = int(task["surah_number"])
    start, end = int(task["from_ayah"]), int(task["to_ayah"])
    log(f"جلب نصوص {task.get('surah_name', '')} من الآية {start} إلى {end}", callback)
    ayahs = get_ayah_texts(surah_number, start, end, tashkeel=bool(style["tashkeel"]), session=session)
    if stop_check and stop_check():
        raise InterruptedError("تم طلب إيقاف التوليد")
    audio_files = []
    durations = []
    for position, _ in enumerate(ayahs, start=1):
        ayah_no = ayahs[position - 1][0]
        log(f"تحميل الصوت: الآية {ayah_no} ({position}/{len(ayahs)})", callback)
        audio = download_audio(surah_number, ayah_no, reciter, session=session, callback=callback)
        audio_files.append(audio)
        durations.append(probe_duration(audio))
        if stop_check and stop_check():
            raise InterruptedError("تم طلب إيقاف التوليد")
    task_slug = safe_name(task.get("id", f"Q-{surah_number}-{start}-{end}"))
    total_duration = sum(durations)
    selected_background = background
    bg_plan: list[dict] | None = None
    bg_mode_used = "single"
    if selected_background is None:
        choices = background_files()
        if choices:
            bg_plan, bg_mode_used = plan_backgrounds(choices, durations, bg_continue=bg_continue,
                                                     mode=bg_mode, callback=callback)
            selected_background = bg_plan[0]["path"]
        else:
            # No background files: fall back to the same generated gradient the
            # preview shows, written once as a still so both paths agree.
            selected_background = TEMP_DIR / "default_background.png"
            if not selected_background.exists():
                default_background().save(selected_background)
            log("لا توجد خلفيات في مجلد «خلفيات» — سيُستخدم تدرّج داكن مولّد تلقائيًا", callback)
    cards_dir = TEMP_DIR / "cards" / task_slug
    segments_dir = TEMP_DIR / "segments" / task_slug
    font_path = ensure_font()
    quran_font_path = ensure_quran_font()
    segments: list[Path] = []
    if bg_plan is None:
        # A background the user pinned by name is one continuous take whatever
        # its length: the explicit choice outranks the planner, which only ever
        # decides among the folder.
        bg_plan = single_background_plan(selected_background, durations, bg_continue=bg_continue)
    for index, ((ayah_no, ayah_text), duration) in enumerate(zip(ayahs, durations), start=1):
        if stop_check and stop_check():
            raise InterruptedError("تم طلب إيقاف التوليد")
        card = cards_dir / f"card_{index:03d}.png"
        render_card(
            ayah_number=ayah_no,
            ayah_text=ayah_text,
            output=card,
            font_path=font_path,
            quran_font_path=font_choice_path(style["font"]) or quran_font_path,
            show_text=bool(style["show_text"]),
            show_number=bool(style["show_number"]),
            font_size=int(style["font_size"]),
            text_color=str(style["text_color"]),
            position=str(style["position"]),
            stroke=bool(style["stroke"]),
            text_bg=bool(style["text_bg"]),
            text_bg_color=str(style["text_bg_color"]),
            text_bg_style=str(style["text_bg_style"]),
            text_bg_opacity=int(style["text_bg_opacity"]),
            text_bg_pad=float(style["text_bg_pad"]),
            text_bg_radius=float(style["text_bg_radius"]),
            shadow=bool(style["shadow"]),
            shadow_opacity=int(style["shadow_opacity"]),
            bg_dim=bg_dim,
        )
        segment = segments_dir / f"segment_{index:03d}.mp4"
        log(f"رندرة الآية {ayah_no} ({index}/{len(ayahs)})", callback)
        entry = bg_plan[index - 1]
        make_segment(entry["path"], audio_files[index - 1], card, segment, duration, animation,
                     bg_offset=entry["offset"], bg_fit=bg_fit)
        segments.append(segment)
    output = OUTPUT_DIR / f"{task_slug}_{safe_name(task.get('surah_name', 'سورة'))}_{start}-{end}.mp4"
    concatenate_segments(segments, output)
    # Only after the file exists: a run that was stopped or failed must not
    # consume footage the user never received a video for.
    if bg_continue:
        # In per-ayah mode the clip draws on several backgrounds, so each one
        # advances by its own share rather than by the whole running time.
        consumed: dict[Path, float] = {}
        for entry, length in zip(bg_plan, durations):
            consumed[entry["path"]] = consumed.get(entry["path"], 0.0) + length
        for used_background, seconds in consumed.items():
            advance_background_cursor(used_background, seconds)
    # Fetched after the video is safely on disk: a tafsir outage should cost a
    # caption section, never the clip itself.
    explanations: list[tuple[int, str]] = []
    tafsir_name = ""
    if tafsir and tafsir != "none":
        tafsir_name = TAFSIR_CHOICES.get(tafsir, TAFSIR_CHOICES["muyassar"])[1]
        log(f"جلب {tafsir_name} للآيات {start}-{end}", callback)
        for ayah_no, _ in ayahs:
            explanation = get_tafsir(surah_number, ayah_no, tafsir, session=session)
            if explanation:
                explanations.append((ayah_no, explanation))
        if not explanations:
            log("تعذّر جلب التفسير — تم إنشاء الوصف بدونه", callback)
    write_social_file(output, task, reciter, topic, tafsir=explanations, tafsir_name=tafsir_name)
    log(f"اكتمل الفيديو: {output.name}", callback)
    return output


def generate_demo(output: Path) -> Path:
    demo_audio = TEMP_DIR / "demo_tone.wav"
    demo_bg = TEMP_DIR / "demo_background.png"
    demo_card = TEMP_DIR / "demo_card.png"
    demo_segment = TEMP_DIR / "demo_segment.mp4"
    TEMP_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.new("RGB", (1080, 1920), (10, 15, 29))
    draw = ImageDraw.Draw(image)
    for y in range(1920):
        ratio = y / 1920
        draw.line((0, y, 1080, y), fill=(int(10 + 15 * ratio), int(15 + 12 * ratio), int(29 + 22 * ratio)))
    image = image.filter(ImageFilter.GaussianBlur(0.6))
    image.save(demo_bg)
    run_ffmpeg(["-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=4", "-c:a", "pcm_s16le", str(demo_audio)])
    render_card(ayah_number=1, ayah_text="بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ", output=demo_card,
                font_path=ensure_font(), quran_font_path=ensure_quran_font(), text_bg=True)
    make_segment(demo_bg, demo_audio, demo_card, demo_segment, 4, "fade")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(demo_segment, output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description="مُصحَف — محرك توليد فيديوهات قرآنية")
    parser.add_argument("--demo", action="store_true", help="إنشاء فيديو اختبار محلي دون اتصال بالإنترنت")
    parser.add_argument("--output", type=Path, default=OUTPUT_DIR / "demo.mp4")
    args = parser.parse_args()
    if args.demo:
        print(f"تم إنشاء فيديو الاختبار: {generate_demo(args.output)}")
        return 0
    parser.error("استخدم لوحة التحكم أو استدعِ generate_task من server.py. وللاختبار استخدم --demo.")
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
