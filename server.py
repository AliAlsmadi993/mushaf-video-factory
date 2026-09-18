from __future__ import annotations

import json
import mimetypes
import os
import subprocess
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from PIL import Image

import drive
from generator import (
    BACKGROUND_DIR,
    BASE_DIR,
    BUNDLE_DIR,
    DEFAULT_STYLE,
    FONT_CHOICES,
    OUTPUT_DIR,
    RECITERS,
    TASKS_XLSX,
    TEMP_DIR,
    background_files,
    background_still,
    ensure_font,
    ensure_quran_font,
    font_choice_path,
    generate_task,
    get_ayah_texts,
    load_tasks,
    render_card,
    background_usage,
    reset_background_cursors,
    TAFSIR_CHOICES,
)

HOST = "0.0.0.0"
PORT = 5055
INDEX_FILE = BASE_DIR / "index.html"
if not INDEX_FILE.exists() and (BUNDLE_DIR / "index.html").exists():
    INDEX_FILE = BUNDLE_DIR / "index.html"

state = {
    "is_running": False,
    "current_index": 0,
    "total_tasks": 0,
    "current_task_name": "",
    "progress_percent": 0,
    "logs": [],
    "should_stop": False,
    "completed_tasks": [],
    "error_tasks": [],
    "started_at": None,
    # filename -> {percent, state, file}. Keyed by the mp4 name so the gallery
    # can match a bar to its card without another request.
    "drive_progress": {},
}
state_lock = threading.Lock()
worker_thread: threading.Thread | None = None


def add_log(message: str) -> None:
    with state_lock:
        state["logs"].append(message)
        state["logs"] = state["logs"][-200:]


def public_state() -> dict:
    with state_lock:
        return dict(state, logs=list(state["logs"]), completed_tasks=list(state["completed_tasks"]),
                    error_tasks=list(state["error_tasks"]), drive_progress=dict(state["drive_progress"]))


def update_workbook_status(task_id: str, status: str) -> None:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(TASKS_XLSX)
        sheet = workbook[workbook.sheetnames[0]]
        header_row = None
        status_col = None
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == "المعرّف":
                    header_row = cell.row
                if cell.value == "حالة التوليد":
                    status_col = cell.column
            if header_row:
                break
        if not header_row:
            workbook.close()
            return
        if not status_col:
            status_col = sheet.max_column + 1
            sheet.cell(header_row, status_col).value = "حالة التوليد"
        id_col = next((cell.column for cell in sheet[header_row] if cell.value == "المعرّف"), None)
        if not id_col:
            workbook.close()
            return
        for row in range(header_row + 1, sheet.max_row + 1):
            if str(sheet.cell(row, id_col).value or "") == task_id:
                sheet.cell(row, status_col).value = status
                break
        workbook.save(TASKS_XLSX)
        workbook.close()
    except Exception as exc:
        add_log(f"تعذر تحديث حالة {task_id} في Excel: {exc}")


def resolve_background(requested: object) -> Path | None:
    """Map a background name from the UI to a file inside BACKGROUND_DIR.

    Returns None for "random"/missing, which lets the caller pick. Names are
    reduced to a bare filename so a request cannot escape the folder.
    """
    name = str(requested or "").strip()
    if not name or name == "random":
        return None
    candidate = (BACKGROUND_DIR / Path(name).name).resolve()
    if candidate.parent == BACKGROUND_DIR.resolve() and candidate.exists():
        return candidate
    return None


def style_from_settings(settings: dict) -> dict:
    """Pick the appearance keys out of a settings payload, ignoring anything
    unexpected so a malformed request can never reshape the output."""
    style = dict(DEFAULT_STYLE)
    for key, default in DEFAULT_STYLE.items():
        if key not in settings or settings[key] is None:
            continue
        value = settings[key]
        try:
            if isinstance(default, bool):
                style[key] = bool(value)
            elif isinstance(default, int):
                style[key] = int(value)
            elif isinstance(default, float):
                style[key] = float(value)
            else:
                style[key] = str(value)
        except (TypeError, ValueError):
            continue
    return style


def bg_fit_from_settings(settings: dict) -> str:
    """Only the two shapes the pipeline knows how to draw; anything else falls
    back to the cover crop so a stray value cannot produce a broken filter."""
    return "blur" if str(settings.get("bg_fit", "")).strip() == "blur" else "cover"


def tafsir_from_settings(settings: dict) -> str:
    choice = str(settings.get("tafsir", "")).strip()
    return choice if choice in TAFSIR_CHOICES else ""


def start_batch(selected_tasks: list[dict], settings: dict) -> None:
    global worker_thread
    with state_lock:
        state.update({
            "is_running": True,
            "current_index": 0,
            "total_tasks": len(selected_tasks),
            "current_task_name": "",
            "progress_percent": 0,
            "logs": [],
            "should_stop": False,
            "completed_tasks": [],
            "error_tasks": [],
            "started_at": time.time(),
            # Bars from the previous batch would otherwise sit there claiming
            # an upload that is no longer happening.
            "drive_progress": {},
        })
    add_log(f"بدأ التوليد لـ {len(selected_tasks)} مقطعاً")
    batch_style = style_from_settings(settings)
    try:
        for index, task in enumerate(selected_tasks, start=1):
            with state_lock:
                requested_stop = bool(state["should_stop"])
                if not requested_stop:
                    state["current_index"] = index
                    state["current_task_name"] = f"{task['id']} — {task['surah_name']}"
                    state["progress_percent"] = int((index - 1) / max(len(selected_tasks), 1) * 100)
            if requested_stop:
                add_log("تم إيقاف العملية بواسطة المستخدم")
                break
            update_workbook_status(task["id"], "جاري التوليد")
            try:
                background = resolve_background(settings.get("background"))
                video = generate_task(
                    task,
                    reciter=settings.get("reciter") or task.get("reciter"),
                    background=background,
                    animation=settings.get("animation", "instant"),
                    # One style dict captured once before the loop drives every
                    # clip in the batch, so they cannot drift apart.
                    style=batch_style,
                    bg_dim=int(settings.get("bg_dim", 0) or 0),
                    topic=settings.get("topic", ""),
                    bg_fit=bg_fit_from_settings(settings),
                    tafsir=tafsir_from_settings(settings),
                    bg_continue=settings.get("bg_continue", True) is not False,
                    bg_mode=str(settings.get("bg_mode") or "auto"),
                    callback=add_log,
                    stop_check=lambda: bool(public_state()["should_stop"]),
                )
                update_workbook_status(task["id"], "مكتمل")
                with state_lock:
                    state["completed_tasks"].append(task["id"])
                # Publishing runs after the file is safely on disk, so a Drive
                # outage costs an upload and never the clip itself.
                if settings.get("drive_upload"):
                    if drive.is_connected():
                        upload_to_drive(video)
                    else:
                        # Skipping in silence is how a whole batch finished with
                        # the upload switch on and nothing to show for it.
                        add_log("تخطّي الرفع إلى Drive: الحساب غير مرتبط — "
                                "أكمل الإعدادات واضغط «ربط حساب Google»")
                    state["progress_percent"] = int(index / max(len(selected_tasks), 1) * 100)
            except InterruptedError:
                update_workbook_status(task["id"], "متوقف")
                add_log("توقف التوليد بأمان")
                break
            except Exception as exc:
                update_workbook_status(task["id"], "فشل")
                with state_lock:
                    state["error_tasks"].append({"id": task["id"], "error": str(exc)})
                add_log(f"فشل {task['id']}: {exc}")
        add_log("انتهت عملية التوليد")
    finally:
        with state_lock:
            state["is_running"] = False
            state["should_stop"] = False
            if state["progress_percent"] < 100 and len(state["completed_tasks"]) == state["total_tasks"]:
                state["progress_percent"] = 100


def json_response(handler: BaseHTTPRequestHandler, payload: object, status: int = 200) -> None:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Content-Length", str(len(body)))
    handler.send_header("Cache-Control", "no-store")
    handler.end_headers()
    handler.wfile.write(body)


def upload_to_drive(video: Path) -> None:
    """Publish one clip, recording progress so the gallery can draw a bar."""
    key = video.name

    def note(payload: dict) -> None:
        with state_lock:
            state["drive_progress"][key] = {
                "percent": payload["percent"],
                "file": payload["file"],
                "state": payload.get("state", "uploading"),
            }

    with state_lock:
        state["drive_progress"][key] = {"percent": 0, "file": key, "state": "uploading"}
    try:
        drive.publish_clip(video, callback=add_log, progress=note)
        with state_lock:
            state["drive_progress"][key] = {"percent": 100, "file": key, "state": "done"}
    except Exception as exc:
        add_log(f"تعذّر الرفع إلى Drive: {exc}")
        with state_lock:
            state["drive_progress"][key] = {"percent": 0, "file": key, "state": "error",
                                            "message": str(exc)[:200]}


def output_items() -> list[dict]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    result = []
    for path in sorted(OUTPUT_DIR.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        if not path.is_file() or path.suffix.lower() != ".mp4":
            continue
        social = path.with_name(path.stem + "_social.txt")
        result.append({
            "filename": path.name,
            "url": f"/output/{path.name}",
            "caption_url": f"/output/{social.name}" if social.exists() else None,
            "size_mb": round(path.stat().st_size / 1024 / 1024, 2),
            "modified": path.stat().st_mtime,
        })
    return result


class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, format, *args):
        return

    def do_GET(self):
        parsed = urlparse(self.path)
        path = unquote(parsed.path)
        if path in ("/", "/index.html"):
            self.serve_file(INDEX_FILE, "text/html; charset=utf-8", no_store=True)
            return
        if path == "/api/tasks":
            try:
                tasks = load_tasks()
                json_response(self, {"tasks": tasks, "count": len(tasks)})
            except Exception as exc:
                json_response(self, {"error": str(exc)}, 500)
            return
        if path == "/api/reciters":
            json_response(self, {"reciters": list(RECITERS.keys())})
            return
        if path == "/api/fonts":
            json_response(self, {"fonts": [{"id": k, "label": v[0]} for k, v in FONT_CHOICES.items()],
                                  "defaults": DEFAULT_STYLE})
            return
        if path == "/api/tafsirs":
            json_response(self, {"tafsirs": [{"id": k, "label": v[1]} for k, v in TAFSIR_CHOICES.items()]})
            return
        if path == "/preview.jpg":
            preview = TEMP_DIR / "preview.jpg"
            if not preview.exists():
                self.send_error(404)
                return
            self.serve_file(preview, "image/jpeg")
            return
        if path == "/api/backgrounds":
            items = []
            for f in background_files():
                usage = background_usage(f)
                items.append({
                    "name": f.name,
                    "type": "video" if usage["is_video"] else "image",
                    "cursor": round(usage["cursor"], 1),
                    "duration": round(usage["duration"], 1),
                    "remaining": round(usage["remaining"], 1),
                })
            json_response(self, {"backgrounds": items})
            return
        if path == "/api/status":
            json_response(self, public_state())
            return
        if path == "/api/drive/status":
            json_response(self, drive.status())
            return
        if path == "/api/outputs":
            json_response(self, {"outputs": output_items()})
            return
        if path.startswith("/output/"):
            name = Path(path[len("/output/"):]).name
            target = (OUTPUT_DIR / name).resolve()
            if target.parent != OUTPUT_DIR.resolve() or not target.exists():
                json_response(self, {"error": "الملف غير موجود"}, 404)
                return
            self.serve_file(target, mimetypes.guess_type(target.name)[0] or "application/octet-stream", ranges=True)
            return
        self.send_error(404)

    def do_HEAD(self):
        path = unquote(urlparse(self.path).path)
        if path.startswith("/output/"):
            name = Path(path[len("/output/"):]).name
            target = (OUTPUT_DIR / name).resolve()
            if target.parent == OUTPUT_DIR.resolve() and target.exists():
                self.send_response(200)
                self.send_header("Content-Type", mimetypes.guess_type(target.name)[0] or "application/octet-stream")
                self.send_header("Content-Length", str(target.stat().st_size))
                self.send_header("Accept-Ranges", "bytes")
                self.end_headers()
                return
        self.send_error(404)

    def do_POST(self):
        path = unquote(urlparse(self.path).path)
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length) if length else b"{}"
        try:
            data = json.loads(raw.decode("utf-8"))
        except Exception:
            json_response(self, {"error": "JSON غير صالح"}, 400)
            return
        if path == "/api/start":
            with state_lock:
                if state["is_running"]:
                    json_response(self, {"error": "يوجد توليد جارٍ بالفعل"}, 409)
                    return
            all_tasks = load_tasks()
            ids = set(str(x) for x in data.get("task_ids", []))
            selected = [task for task in all_tasks if task["id"] in ids]
            custom = data.get("custom_task")
            if custom:
                selected.insert(0, {
                    "index": 0,
                    "id": "CUSTOM",
                    "juz": 0,
                    "surah_number": int(custom.get("surah_number", 1)),
                    "surah_name": str(custom.get("surah_name", "سورة مختارة")),
                    "from_ayah": int(custom.get("from_ayah", 1)),
                    "to_ayah": int(custom.get("to_ayah", 1)),
                    "ayah_count": int(custom.get("to_ayah", 1)) - int(custom.get("from_ayah", 1)) + 1,
                    "opening": "",
                    "reciter": str(data.get("reciter") or "ياسر الدوسري"),
                })
            if not selected:
                json_response(self, {"error": "اختر مقطعاً واحداً على الأقل"}, 400)
                return
            worker_thread = threading.Thread(target=start_batch, args=(selected, data), daemon=True, name="mushaf-generator")
            worker_thread.start()
            json_response(self, {"ok": True, "selected": len(selected)})
            return
        if path == "/api/preview":
            try:
                surah = int(data.get("surah_number") or 1)
                ayah = int(data.get("ayah") or 1)
                style = style_from_settings(data)
                texts = get_ayah_texts(surah, ayah, ayah, tashkeel=bool(style["tashkeel"]))
                preview = TEMP_DIR / "preview.jpg"
                card_file = TEMP_DIR / "preview_card.png"
                render_card(
                    ayah_number=texts[0][0],
                    ayah_text=texts[0][1],
                    output=card_file,
                    font_path=ensure_font(),
                    quran_font_path=font_choice_path(style["font"]) or ensure_quran_font(),
                    show_text=bool(style["show_text"]),
                    show_number=bool(style["show_number"]),
                    font_size=int(style["font_size"]),
                    text_color=str(style["text_color"]),
                    position=str(style["position"]),
                    stroke=bool(style["stroke"]),
                    text_bg=bool(style["text_bg"]),
                    text_bg_color=str(style["text_bg_color"]),
                    text_bg_style=str(style["text_bg_style"]),
                    text_bg_opacity=int(style["text_bg_opacity"]),
                    text_bg_pad=float(style["text_bg_pad"]),
                    text_bg_radius=float(style["text_bg_radius"]),
                    shadow=bool(style["shadow"]),
                    shadow_opacity=int(style["shadow_opacity"]),
                    bg_dim=int(data.get("bg_dim", 0) or 0),
                )
                # Composite over the very background the clip will use, so the
                # preview is a true first frame rather than the bare card. The
                # dim lives in the card's alpha, so the still is undimmed here.
                # On "random" the batch picks a different file per clip; the
                # preview settles on the first available one so it shows a real
                # background instead of the empty-folder gradient.
                chosen = resolve_background(data.get("background"))
                if chosen is None:
                    available = background_files()
                    chosen = available[0] if available else None
                still = background_still(chosen, bg_fit=bg_fit_from_settings(data)).convert("RGBA")
                still.alpha_composite(Image.open(card_file).convert("RGBA"))
                still.convert("RGB").save(preview, "JPEG", quality=88)
                json_response(self, {"ok": True, "url": f"/preview.jpg?t={int(time.time()*1000)}"})
            except Exception as exc:
                json_response(self, {"error": str(exc)}, 500)
            return
        if path == "/api/drive/config":
            try:
                config = drive.save_config(
                    client_id=data.get("client_id"),
                    client_secret=data.get("client_secret"),
                    root_folder=data.get("root_folder"),
                    auto_upload=data.get("auto_upload"),
                )
                missing = [label for key, label in (("client_id", "معرّف العميل"),
                                                    ("client_secret", "سر العميل"))
                           if not config[key]]
                if missing:
                    add_log("إعدادات Drive ناقصة: " + " و".join(missing))
                else:
                    add_log(f"حُفظت إعدادات Drive — المجلد الرئيسي «{config['root_folder']}»")
                json_response(self, drive.status())
            except Exception as exc:
                json_response(self, {"error": str(exc)}, 400)
            return
        if path == "/api/drive/connect":
            # The consent flow waits on a browser, so it must not block the
            # HTTP handler; the page polls /api/drive/status for the result.
            def run_connect():
                try:
                    drive.connect(callback=add_log)
                except Exception as exc:
                    add_log(f"تعذّر ربط Drive: {exc}")
            threading.Thread(target=run_connect, daemon=True, name="drive-connect").start()
            add_log("جارٍ فتح صفحة تسجيل الدخول إلى Google…")
            json_response(self, {"ok": True})
            return
        if path == "/api/drive/disconnect":
            drive.disconnect()
            add_log("تم فصل حساب Google Drive")
            json_response(self, drive.status())
            return
        if path == "/api/drive/upload":
            name = Path(str(data.get("filename", ""))).name
            target = (OUTPUT_DIR / name).resolve()
            if target.parent != OUTPUT_DIR.resolve() or not target.exists():
                json_response(self, {"error": "الملف غير موجود"}, 404)
                return

            threading.Thread(target=upload_to_drive, args=(target,), daemon=True,
                             name="drive-upload").start()
            json_response(self, {"ok": True})
            return
        if path == "/api/reset_backgrounds":
            reset_background_cursors()
            add_log("تم تصفير مؤشرات الخلفيات — ستبدأ كل خلفية من أولها")
            json_response(self, {"ok": True})
            return
        if path == "/api/stop":
            with state_lock:
                state["should_stop"] = True
            add_log("تم تسجيل طلب الإيقاف…")
            json_response(self, {"ok": True})
            return
        if path == "/api/open_folder":
            folder_name = str(data.get("folder", "output"))
            folder = BACKGROUND_DIR if folder_name in {"backgrounds", "خلفيات"} else OUTPUT_DIR
            folder.mkdir(parents=True, exist_ok=True)
            try:
                if os.name == "nt":
                    os.startfile(str(folder))
                elif sys_platform() == "darwin":
                    subprocess.Popen(["open", str(folder)])
                else:
                    subprocess.Popen(["xdg-open", str(folder)])
            except Exception as exc:
                json_response(self, {"error": str(exc)}, 500)
                return
            json_response(self, {"ok": True, "path": str(folder)})
            return
        self.send_error(404)

    def serve_file(self, path: Path, content_type: str, ranges: bool = False, no_store: bool = False):
        if not path.exists():
            self.send_error(404)
            return
        total = path.stat().st_size
        start, end = 0, total - 1
        status = 200
        if ranges and self.headers.get("Range"):
            value = self.headers["Range"].replace("bytes=", "")
            first, _, last = value.partition("-")
            try:
                start = int(first) if first else 0
                end = int(last) if last else total - 1
                start = max(0, start)
                end = min(total - 1, end)
                if start > end:
                    raise ValueError
                status = 206
            except ValueError:
                self.send_error(416)
                return
        length = end - start + 1
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(length))
        if no_store:
            # Without this a rebuilt interface stays hidden behind the copy the
            # browser cached from an earlier run.
            self.send_header("Cache-Control", "no-store, must-revalidate")
        if ranges:
            self.send_header("Accept-Ranges", "bytes")
            if status == 206:
                self.send_header("Content-Range", f"bytes {start}-{end}/{total}")
        self.end_headers()
        with path.open("rb") as handle:
            handle.seek(start)
            remaining = length
            while remaining:
                chunk = handle.read(min(1024 * 256, remaining))
                if not chunk:
                    break
                self.wfile.write(chunk)
                remaining -= len(chunk)


def sys_platform() -> str:
    import sys
    return sys.platform


def main():
    BACKGROUND_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"مُصحَف يعمل على http://{HOST}:{PORT}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nتم إيقاف الخادم")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
